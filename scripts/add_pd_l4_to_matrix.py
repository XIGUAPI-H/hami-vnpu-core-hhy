#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""在 v3 基础上补充 PD 分离相关 L4 到模型推理域。"""
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

INPUT = Path(r"C:\Users\d00804096\Desktop\AI平台功能设计矩阵_v3.xlsx")
OUTPUT = Path(r"C:\Users\d00804096\Desktop\AI平台功能设计矩阵_v4.xlsx")
SOURCE_XLSX = Path(r"C:\Users\d00804096\Desktop\功能项分析2026.xlsx")
PLATFORMS = ["招行", "openfuyao", "MM", "MA", "MindCluster"]

# 新增：模型推理 → PD分离部署
PD_ITEMS = [
    {
        "L1": "模型推理",
        "L2": "推理服务部署",
        "L3": "分布式部署",
        "L4": "PD分离多角色部署",
        "功能描述": (
            "将推理拆成 Prefill（处理输入 prompt）与 Decode（逐 token 生成）两类角色，"
            "分别部署为独立服务或实例池；支持多机多卡下按 P/D 角色分配算力，"
            "是大模型高并发在线推理的常见架构（可与张量并行组合使用）。"
        ),
        "kw": ["PD分离部署", "PD 分离", "多角色部署", "Prefill", "Decode", "大EP"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "Prefill服务",
        "L4": "Prefill节点部署",
        "功能描述": (
            "单独部署 Prefill 服务：负责接收用户 prompt 并完成首段计算，"
            "产出 KV Cache 等中间状态供 Decode 使用；可按 prompt 长度与并发单独扩缩。"
        ),
        "kw": ["Prefill", "PD分离", "P节点", "prefill"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "Decode服务",
        "L4": "Decode节点部署",
        "功能描述": (
            "单独部署 Decode 服务：接收 Prefill 传递的中间状态，逐 token 生成输出；"
            "通常副本数多于 Prefill，以应对生成阶段的高并发与低时延要求。"
        ),
        "kw": ["Decode", "PD分离", "D节点", "decode"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "弹性伸缩",
        "L4": "PD分离动态扩缩",
        "功能描述": (
            "Prefill 池与 Decode 池分别监控负载（如队列长度、token 吞吐），"
            "独立触发 HPA 或自定义扩缩策略，避免「整服务一起扩缩」造成的资源浪费。"
        ),
        "kw": ["PD 分离动态扩缩", "PD分离动态扩缩", "动态扩缩", "PD分离扩缩"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "推理优化",
        "L4": "PD分离推理优化",
        "功能描述": (
            "针对 P/D 分离架构做推理链路优化，如 P→D 状态传递、批处理策略、"
            "长 prompt 与短 prompt 分流等，降低端到端时延、提升吞吐。"
        ),
        "kw": ["PD 分离推理优化", "PD分离推理优化", "推理优化"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "KV缓存",
        "L4": "分布式KVCache",
        "功能描述": (
            "在 P/D 节点或跨机之间共享/传递 KV Cache（如 Mooncake 方案），"
            "减少重复计算与显存占用，提升多轮对话与高并发场景下的性能。"
        ),
        "kw": ["Mooncake", "KVCache", "KV Cache", "分布式 KVCache"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "引擎适配",
        "L4": "AIBrix vLLM PD分离",
        "功能描述": (
            "基于 AIBrix + vLLM 引擎的 PD 分离部署与调度适配，"
            "在 K8s 环境下按 P/D 角色创建与管理推理实例。"
        ),
        "kw": ["AIBrix", "vLLM PD", "PD 分离"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "引擎适配",
        "L4": "OME SGLang PD分离",
        "功能描述": (
            "基于 OME + SGLang 引擎的 PD 分离部署方案，"
            "支持 SGLang 运行时下的 Prefill/Decode 分池与扩缩。"
        ),
        "kw": ["OME", "SGLang PD", "SGLang", "PD 分离"],
        "阶段": "阶段2",
    },
    {
        "L1": "模型推理",
        "L2": "PD分离部署",
        "L3": "路由调度",
        "L4": "PD请求路由与调度",
        "功能描述": (
            "在推理网关或调度层识别请求阶段，将流量路由到 Prefill 或 Decode 池，"
            "并协调 P→D 的状态 handoff，对调用方仍保持统一 API 入口。"
        ),
        "kw": ["PD", "路由", "调度", "handoff", "网关"],
        "阶段": "阶段2",
    },
]


def norm(s: str) -> str:
    return re.sub(r"[\s\-_/（）()]", "", str(s).strip().lower())


def load_platform(path: Path, sheet: str) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet, header=1)
    cols = list(df.columns)
    rename = {cols[0]: "L1", cols[1]: "L2", cols[2]: "L3", cols[3]: "L4"}
    if len(cols) > 4:
        rename[cols[4]] = "功能描述"
    df = df.rename(columns=rename)
    for c in ["L1", "L2", "L3"]:
        if c in df.columns:
            df[c] = df[c].ffill()
    df = df[df["L4"].notna()].copy()
    df["功能描述"] = df.get("功能描述", pd.Series([""] * len(df))).fillna("")
    return df


def score_row(kw_list: list[str], row: pd.Series) -> int:
    text = norm(f"{row['L1']}{row['L2']}{row['L3']}{row['L4']}{row['功能描述']}")
    l4 = norm(row["L4"])
    score = 0
    for kw in kw_list:
        k = norm(kw)
        if not k:
            continue
        if k == l4:
            score += 20
        elif k in l4 or l4 in k:
            score += 12
        elif k in text:
            score += 4
    return score


def match_platforms(kw: list[str], platform_data: dict) -> dict[str, str]:
    out = {}
    for plat in PLATFORMS:
        matches = []
        for _, row in platform_data[plat].iterrows():
            s = score_row(kw, row)
            if s >= 10:
                matches.append((s, str(row["L4"]).strip()))
        matches.sort(key=lambda x: (-x[0], x[1]))
        seen, picked = set(), []
        for _, l4 in matches:
            if l4 in seen:
                continue
            seen.add(l4)
            picked.append(l4)
            if len(picked) >= 3:
                break
        out[plat] = "；".join(picked)
    return out


def estimate_workload(desc: str, l4: str, matched: str) -> tuple[float, float]:
    text = (desc or "") + (l4 or "") + (matched or "")
    if any(k in text for k in ("分布式", "PD分离", "Mooncake", "多角色", "扩缩")):
        dev = 2.0
    elif any(k in text for k in ("部署", "路由", "引擎")):
        dev = 1.0
    else:
        dev = 0.5
    if not matched:
        dev = round(dev * 1.2, 2)
    test = round(dev * 0.4, 2)
    return dev, test


def phase_hint(phase: str, l2: str, l4: str) -> str:
    if phase == "阶段1":
        return f"阶段①轻量核心：{l2}——{l4}（PoC 必备）"
    if phase == "阶段3":
        return f"阶段③高阶能力：PD分离——{l4}"
    return f"阶段②规模化：{l2}——{l4}（PD分离/增强）"


def blank_hierarchy(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col, parents in [("L1", []), ("L2", ["L1"]), ("L3", ["L1", "L2"])]:
        prev = None
        for i in out.index:
            key = tuple(str(out.at[i, p]) for p in parents + [col])
            if key == prev:
                out.at[i, col] = ""
            else:
                prev = key
    return out


def merge_cells(ws, df: pd.DataFrame, header_row: int = 1) -> None:
    n = len(df)
    start = header_row + 1
    if n <= 1:
        return

    def merge_col(letter: str, keys: pd.Series) -> None:
        i = 0
        while i < n:
            j = i + 1
            while j < n and keys.iloc[j] == keys.iloc[i]:
                j += 1
            if j - i > 1:
                ws.merge_cells(f"{letter}{start + i}:{letter}{start + j - 1}")
            i = j

    merge_col("A", df["L1"].astype(str))
    merge_col("B", df["L1"].astype(str) + "\0" + df["L2"].astype(str))
    merge_col(
        "C",
        df["L1"].astype(str) + "\0" + df["L2"].astype(str) + "\0" + df["L3"].astype(str),
    )


def style_sheet(ws, header_row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="B4C7E7")
    font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="B4B4B4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(header_row, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = center
        cell.border = border
    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = left if c == 5 else center
    for letter, w in {
        "A": 12, "B": 14, "C": 14, "D": 24, "E": 56,
        "F": 18, "G": 18, "H": 18, "I": 18, "J": 18,
        "K": 10, "L": 10, "M": 10, "N": 8, "O": 32,
    }.items():
        ws.column_dimensions[letter].width = w


def main():
    df = pd.read_excel(INPUT, sheet_name=0)
    platform_data = {p: load_platform(SOURCE_XLSX, p) for p in PLATFORMS}

    # 去掉算力调度里重复的「推理PD分离扩缩」，改由模型推理 PD分离部署 承载
    before = len(df)
    df = df[~df["L4"].astype(str).str.contains("推理PD分离扩缩", na=False)].copy()
    removed = before - len(df)

    new_rows = []
    for item in PD_ITEMS:
        plat = match_platforms(item["kw"], platform_data)
        matched = " ".join(plat.values())
        dev, test = estimate_workload(item["功能描述"], item["L4"], matched)
        cover = sum(1 for v in plat.values() if v)
        new_rows.append(
            {
                **{k: item[k] for k in ["L1", "L2", "L3", "L4", "功能描述", "阶段"]},
                **{f"{p}对应": plat[p] for p in PLATFORMS},
                "平台覆盖数": cover,
                "开发工作量(人月)": dev,
                "测试工作量(人月)": test,
                "总工作量(人月)": round(dev + test, 2),
                "阶段说明": phase_hint(item["阶段"], item["L2"], item["L4"]),
            }
        )
    new_df = pd.DataFrame(new_rows)

    # 插入位置：模型推理 → 推理服务部署 块之后（API服务 之前）
    for c in ["L1", "L2", "L3"]:
        df[c] = df[c].ffill()
    insert_at = None
    for i, r in df.iterrows():
        if r["L1"] == "模型推理" and r["L2"] == "API服务":
            insert_at = i
            break
    if insert_at is None:
        # fallback: after last 模型推理 row
        idx = df.index[df["L1"].ffill() == "模型推理"]
        insert_at = idx[-1] + 1 if len(idx) else len(df)

    top = df.iloc[:insert_at]
    bottom = df.iloc[insert_at:]
    out = pd.concat([top, new_df, bottom], ignore_index=True)

    # 更新阶段汇总
    summary = []
    for p in ["阶段1", "阶段2", "阶段3"]:
        sub = out[out["阶段"] == p]
        summary.append(
            {
                "阶段": p,
                "功能点数": len(sub),
                "开发工作量(人月)": round(sub["开发工作量(人月)"].sum(), 1),
                "测试工作量(人月)": round(sub["测试工作量(人月)"].sum(), 1),
                "总工作量(人月)": round(sub["总工作量(人月)"].sum(), 1),
            }
        )

    xl = pd.ExcelFile(INPUT)
    other_sheets = {n: pd.read_excel(INPUT, sheet_name=n) for n in xl.sheet_names if n not in ("功能设计总表", "阶段汇总")}

    display = blank_hierarchy(out)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        display.to_excel(writer, sheet_name="功能设计总表", index=False)
        ws = writer.sheets["功能设计总表"]
        style_sheet(ws)
        merge_cells(ws, out)
        pd.DataFrame(summary).to_excel(writer, sheet_name="阶段汇总", index=False)
        style_sheet(writer.sheets["阶段汇总"])
        for name, sdf in other_sheets.items():
            sdf.to_excel(writer, sheet_name=name, index=False)

    print(f"Written: {OUTPUT}")
    print(f"Removed from 算力调度: {removed} row(s)")
    print(f"Added PD L4: {len(new_df)}")
    print(f"Total rows: {len(out)}")
    print(pd.DataFrame(summary).to_string(index=False))
    print("\nNew L4:")
    for r in new_rows:
        print(f"  - {r['L4']} | 覆盖{r['平台覆盖数']}平台")


if __name__ == "__main__":
    main()
