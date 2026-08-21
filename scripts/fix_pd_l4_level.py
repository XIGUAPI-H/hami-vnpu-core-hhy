#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""修正：PD 分离只在 L4，L2 归入「推理服务部署」。"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

INPUT = Path(r"C:\Users\d00804096\Desktop\AI平台功能设计矩阵_v4.xlsx")
OUTPUT = Path(r"C:\Users\d00804096\Desktop\AI平台功能设计矩阵_v4.xlsx")

# L4 -> L3（L2 统一为「推理服务部署」）
L3_MAP = {
    "PD分离多角色部署": "分布式部署",
    "Prefill节点部署": "分布式部署",
    "Decode节点部署": "分布式部署",
    "PD分离动态扩缩": "弹性伸缩",
    "PD分离推理优化": "推理优化",
    "分布式KVCache": "KV缓存",
    "AIBrix vLLM PD分离": "引擎适配",
    "OME SGLang PD分离": "引擎适配",
    "PD请求路由与调度": "服务路由",
}

PD_L4 = set(L3_MAP.keys())


def blank_hierarchy(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col, parents in [("L1", []), ("L2", ["L1"]), ("L3", ["L1", "L2"])]:
        prev = None
        for i in out.index:
            key = tuple(str(out.at[i, p]) for p in parents + [col])
            if key == prev:
                out.at[i, col] = ""
            else:
                prev = key
    return out


def merge_cells(ws, df: pd.DataFrame, header_row: int = 1) -> None:
    n = len(df)
    start = header_row + 1
    if n <= 1:
        return

    def merge_col(letter: str, keys: pd.Series) -> None:
        i = 0
        while i < n:
            j = i + 1
            while j < n and keys.iloc[j] == keys.iloc[i]:
                j += 1
            if j - i > 1:
                ws.merge_cells(f"{letter}{start + i}:{letter}{start + j - 1}")
            i = j

    merge_col("A", df["L1"].astype(str))
    merge_col("B", df["L1"].astype(str) + "\0" + df["L2"].astype(str))
    merge_col(
        "C",
        df["L1"].astype(str) + "\0" + df["L2"].astype(str) + "\0" + df["L3"].astype(str),
    )


def style_sheet(ws, header_row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="B4C7E7")
    font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="B4B4B4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(header_row, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = center
        cell.border = border
    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = left if c == 5 else center


def main():
    xl = pd.ExcelFile(INPUT)
    df = pd.read_excel(INPUT, sheet_name=0)
    other = {n: pd.read_excel(INPUT, sheet_name=n) for n in xl.sheet_names if n != "功能设计总表"}

    fixed = 0
    for i, row in df.iterrows():
        l4 = str(row["L4"]).strip()
        if l4 not in PD_L4 and row.get("L2") != "PD分离部署":
            continue
        if l4 in PD_L4:
            df.at[i, "L1"] = "模型推理"
            df.at[i, "L2"] = "推理服务部署"
            df.at[i, "L3"] = L3_MAP[l4]
            hint = str(row.get("阶段说明", ""))
            if "PD分离部署" in hint:
                df.at[i, "阶段说明"] = hint.replace("PD分离部署", "推理服务部署")
            fixed += 1

    # 插入位置：PD 相关 L4 紧跟「推理弹性扩缩容」之后
    for c in ["L1", "L2", "L3", "L4"]:
        df[c] = df[c].astype(str).replace("nan", "")

    pd_rows = df[df["L4"].isin(PD_L4)].copy()
    rest = df[~df["L4"].isin(PD_L4)].copy()
    for c in ["L1", "L2", "L3"]:
        rest[c] = rest[c].replace("", pd.NA).ffill()

    insert_at = None
    for i, r in rest.iterrows():
        if str(r["L4"]).strip() == "推理弹性扩缩容":
            insert_at = rest.index.get_loc(i) + 1
            break
    if insert_at is None:
        for i, r in rest.iterrows():
            if str(r.get("L2", "")).strip() == "API服务":
                insert_at = rest.index.get_loc(i)
                break
    if insert_at is None:
        merged = pd.concat([rest, pd_rows], ignore_index=True)
    else:
        top = rest.iloc[:insert_at]
        bottom = rest.iloc[insert_at:]
        merged = pd.concat([top, pd_rows, bottom], ignore_index=True)

    # 阶段汇总
    summary = []
    for p in ["阶段1", "阶段2", "阶段3"]:
        sub = merged[merged["阶段"] == p]
        summary.append(
            {
                "阶段": p,
                "功能点数": len(sub),
                "开发工作量(人月)": round(sub["开发工作量(人月)"].sum(), 1),
                "测试工作量(人月)": round(sub["测试工作量(人月)"].sum(), 1),
                "总工作量(人月)": round(sub["总工作量(人月)"].sum(), 1),
            }
        )

    display = blank_hierarchy(merged)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        display.to_excel(writer, sheet_name="功能设计总表", index=False)
        ws = writer.sheets["功能设计总表"]
        style_sheet(ws)
        merge_cells(ws, merged)
        pd.DataFrame(summary).to_excel(writer, sheet_name="阶段汇总", index=False)
        style_sheet(writer.sheets["阶段汇总"])
        for name, sdf in other.items():
            if name != "阶段汇总":
                sdf.to_excel(writer, sheet_name=name, index=False)

    print(f"Fixed {fixed} rows: L2=推理服务部署, PD only in L4")
    print(f"Written: {OUTPUT}")
    print("\nPD L4 under 推理服务部署:")
    for l4, l3 in L3_MAP.items():
        print(f"  L3 {l3} → L4 {l4}")


if __name__ == "__main__":
    main()
