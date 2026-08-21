#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按轻量核心/重量级重新划分阶段，更新设计矩阵。"""
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

DESKTOP = Path(r"C:\Users\d00804096\Desktop")
INPUT = DESKTOP / "AI平台功能设计矩阵_v2.xlsx"
OUTPUT = DESKTOP / "AI平台功能设计矩阵_v3.xlsx"

PHASE_FINISH = {"阶段1": "2026年12月", "阶段2": "2027年6月", "阶段3": "2027年12月"}

# L4 -> 阶段（轻量核心=1，规模化/集群=2，高阶AI=3）
PHASE_BY_L4: dict[str, int] = {
    # --- 阶段1：轻量核心（单业务 PoC：训-推-管-用跑通）---
    "在线推理服务部署": 1,
    "模型广场体验": 1,
    "在线对话": 1,
    "对话推理API": 1,
    "向量化API": 1,
    "Function Call": 1,
    "知识库创建与管理": 1,
    "RAG检索问答": 1,
    "重排与召回": 1,
    "推理调用量统计": 1,
    "LoRA精调": 1,
    "精调作业提交": 1,
    "Notebook交互开发": 1,
    "训练作业监控": 1,
    "训练数据集管理": 1,
    "模型注册与纳管": 1,
    "模型版本管理": 1,
    "模型广场与订阅": 1,
    "数据集全生命周期": 1,
    "镜像仓库管理": 1,
    "知识库文档管理": 1,
    "应用与API Key": 1,
    "工作空间管理": 1,
    "子系统与租户边界": 1,
    "租户生命周期": 1,
    "租户配额管理": 1,
    "命名空间隔离": 1,
    "API Key管理": 1,
    "用户与角色管理": 1,
    "成员与审批": 1,
    "NPU/GPU利用率监控": 1,
    "推理服务监控": 1,
    "训练任务监控": 1,
    "异常告警": 1,
    "日志采集与查询": 1,
    # --- 阶段3：高阶 AI / 自治运维 ---
    "多模态推理API": 3,
    "RLHF训练": 3,
    "DPO/偏好对齐": 3,
    "Agentic RL": 3,
    "节点故障感知": 3,
}
# 未列出的 L4 默认阶段2（规模化、集群、高级训推、完整运维）


def phase_hint(phase: int, l1: str, l2: str, l4: str) -> str:
    if phase == 1:
        return f"阶段①轻量核心：{l2}——{l4}（PoC 必备）"
    if phase == 3:
        return f"阶段③高阶能力：{l1}——{l4}"
    return f"阶段②规模化：{l2}——{l4}（集群/增强）"


def narrative(phase: int, sub: pd.DataFrame) -> str:
    cnt = len(sub)
    total = round(sub["总工作量(人月)"].sum(), 1)
    finish = PHASE_FINISH[f"阶段{phase}"]
    l1s = "、".join(sub["L1"].dropna().unique()[:5])
    if phase == 1:
        return (
            f"阶段①轻量核心，覆盖训-推-管-用主链路（{l1s} 等），"
            f"支撑单业务 PoC 上线；共 {cnt} 项，预计 {total} 人月，{finish} 完成"
        )
    if phase == 2:
        return (
            f"阶段②规模化增强，在轻量核心之上补齐分布式训推、算力调度、集群运维与成本治理（{l1s} 等），"
            f"共 {cnt} 项，预计 {total} 人月，{finish} 完成"
        )
    return (
        f"阶段③高阶能力，补齐多模态、强化学习对齐与智能故障感知等（{l1s}），"
        f"共 {cnt} 项，预计 {total} 人月，{finish} 完成"
    )


def main():
    shutil.copy2(INPUT, OUTPUT)

    df = pd.read_excel(OUTPUT, sheet_name=0)
    for c in ["L1", "L2", "L3"]:
        df[c] = df[c].ffill()

    phases = []
    hints = []
    for _, r in df.iterrows():
        l4 = str(r["L4"]).strip()
        p = PHASE_BY_L4.get(l4, 2)
        phases.append(f"阶段{p}")
        hints.append(phase_hint(p, str(r["L1"]), str(r["L2"]), l4))

    df["阶段"] = phases
    df["阶段说明"] = hints

    # 用 openpyxl 只更新阶段列，保留合并单元格
    wb = load_workbook(OUTPUT)
    ws = wb.worksheets[0]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    ci, hi = headers["阶段"], headers["阶段说明"]
    l1 = l2 = l3 = None
    ri = 0
    for r in range(2, ws.max_row + 1):
        v4 = ws.cell(r, headers["L4"]).value
        if v4 in (None, ""):
            continue
        ws.cell(r, ci, phases[ri])
        cell = ws.cell(r, hi, hints[ri])
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ri += 1
    wb.save(OUTPUT)

    # 阶段汇总 sheet
    summary_rows = []
    narratives = {}
    for p in [1, 2, 3]:
        sub = df[df["阶段"] == f"阶段{p}"]
        narratives[p] = narrative(p, sub)
        summary_rows.append(
            {
                "阶段": f"阶段{p}",
                "阶段目标": narratives[p],
                "功能点数": len(sub),
                "开发工作量(人月)": round(sub["开发工作量(人月)"].sum(), 1),
                "测试工作量(人月)": round(sub["测试工作量(人月)"].sum(), 1),
                "总工作量(人月)": round(sub["总工作量(人月)"].sum(), 1),
                "预计完成": PHASE_FINISH[f"阶段{p}"],
            }
        )
    summary_df = pd.DataFrame(summary_rows)

    with pd.ExcelWriter(OUTPUT, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
        summary_df.to_excel(w, sheet_name="阶段汇总", index=False)

    print(f"Write: {OUTPUT}")
    print(summary_df.to_string(index=False))
    for p in [1, 2, 3]:
        print(f"\n--- 阶段{p} ({len(df[df['阶段']==f'阶段{p}'])}) ---")
        print(narratives[p])


if __name__ == "__main__":
    main()
