#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""多机多卡部署 L3 + PD 分离 L4 + 加宽列格式。"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

V4 = Path(r"C:\Users\d00804096\Desktop\AI平台功能设计矩阵_v4.xlsx")
OUTPUT = V4

# 加宽列宽，确保功能描述与各平台对应列能完整显示
COL_WIDTHS = {
    "A": 14,   # L1
    "B": 16,   # L2
    "C": 18,   # L3
    "D": 26,   # L4
    "E": 72,   # 功能描述
    "F": 22,   # 招行
    "G": 22,   # openfuyao
    "H": 22,   # MM
    "I": 22,   # MA
    "J": 24,   # MindCluster
    "K": 12,   # 平台覆盖数
    "L": 14,   # 开发
    "M": 14,   # 测试
    "N": 14,   # 总工作量
    "O": 10,   # 阶段
    "P": 36,   # 阶段说明
}

PD_L4 = {
    "PD分离多角色部署",
    "Prefill节点部署",
    "Decode节点部署",
    "PD分离动态扩缩",
    "PD分离推理优化",
    "分布式KVCache",
    "AIBrix vLLM PD分离",
    "OME SGLang PD分离",
    "PD请求路由与调度",
}

# PD 相关 L4 归入 L3「多机多卡部署」
PD_L3 = "多机多卡部署"


def blank_hierarchy(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col, parents in [("L1", []), ("L2", ["L1"]), ("L3", ["L1", "L2"])]:
        prev = None
        for i in out.index:
            key = tuple(str(out.at[i, p]) for p in parents + [col])
            if key == prev:
                out.at[i, col] = ""
            else:
                prev = key
    return out


def merge_cells(ws, df: pd.DataFrame, header_row: int = 1) -> None:
    n = len(df)
    start = header_row + 1
    if n <= 1:
        return

    def merge_col(letter: str, keys: pd.Series) -> None:
        i = 0
        while i < n:
            j = i + 1
            while j < n and keys.iloc[j] == keys.iloc[i]:
                j += 1
            if j - i > 1:
                ws.merge_cells(f"{letter}{start + i}:{letter}{start + j - 1}")
            i = j

    merge_col("A", df["L1"].astype(str))
    merge_col("B", df["L1"].astype(str) + "\0" + df["L2"].astype(str))
    merge_col(
        "C",
        df["L1"].astype(str) + "\0" + df["L2"].astype(str) + "\0" + df["L3"].astype(str),
    )


def apply_styles(ws, header_row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="B4C7E7")
    font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin = Side(style="thin", color="B4B4B4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = {ws.cell(header_row, c).value: c for c in range(1, ws.max_column + 1)}
    desc_col = headers.get("功能描述")
    stage_col = headers.get("阶段说明")

    for c in range(1, ws.max_column + 1):
        cell = ws.cell(header_row, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = center
        cell.border = border

    for r in range(header_row + 1, ws.max_row + 1):
        max_lines = 2
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.border = border
            if c == desc_col or c == stage_col:
                cell.alignment = left
                val = str(cell.value or "")
                max_lines = max(max_lines, val.count("\n") + 1, len(val) // 40 + 1)
            elif c in (headers.get("招行对应"), headers.get("openfuyao对应"),
                       headers.get("MM对应"), headers.get("MA对应"), headers.get("MindCluster对应")):
                cell.alignment = left
            else:
                cell.alignment = center
        ws.row_dimensions[r].height = min(120, max(36, max_lines * 15))

    for letter, w in COL_WIDTHS.items():
        ws.column_dimensions[letter].width = w
    for col in range(1, ws.max_column + 1):
        letter = ws.cell(header_row, col).column_letter
        if letter not in COL_WIDTHS:
            ws.column_dimensions[letter].width = 18


def reorder_inference_deploy(df: pd.DataFrame) -> pd.DataFrame:
    for c in ["L1", "L2", "L3"]:
        df[c] = df[c].ffill()

    # L3 分布式部署 → 多机多卡部署
    df.loc[df["L3"].astype(str).str.strip() == "分布式部署", "L3"] = PD_L3

    # 多机多卡分布式推理 也归入 多机多卡部署 L3
    mask_infer = df["L4"].astype(str).str.strip() == "多机多卡分布式推理"
    df.loc[mask_infer, "L3"] = PD_L3

    # 所有 PD 相关 L4 统一 L3 = 多机多卡部署
    for i, row in df.iterrows():
        l4 = str(row["L4"]).strip()
        if l4 in PD_L4 or l4 == "多机多卡分布式推理":
            df.at[i, "L1"] = "模型推理"
            df.at[i, "L2"] = "推理服务部署"
            df.at[i, "L3"] = PD_L3

    # 重排：多机多卡块 = 多机多卡分布式推理 + PD 各 L4，放在推理弹性扩缩容之前
    multi_l4_order = [
        "多机多卡分布式推理",
        "PD分离多角色部署",
        "Prefill节点部署",
        "Decode节点部署",
        "PD分离动态扩缩",
        "PD分离推理优化",
        "分布式KVCache",
        "AIBrix vLLM PD分离",
        "OME SGLang PD分离",
        "PD请求路由与调度",
    ]
    multi_rows = []
    rest_rows = []
    for _, row in df.iterrows():
        l4 = str(row["L4"]).strip()
        if l4 in multi_l4_order:
            multi_rows.append(row)
        else:
            rest_rows.append(row)
    order_map = {name: i for i, name in enumerate(multi_l4_order)}
    multi_rows.sort(key=lambda r: order_map.get(str(r["L4"]).strip(), 999))
    multi_df = pd.DataFrame(multi_rows)
    rest_df = pd.DataFrame(rest_rows)

    insert_at = None
    for i, r in rest_df.iterrows():
        if str(r["L4"]).strip() == "推理弹性扩缩容":
            insert_at = rest_df.index.get_loc(i)
            break
    if insert_at is None:
        for i, r in rest_df.iterrows():
            if str(r.get("L2", "")).strip() == "API服务":
                insert_at = rest_df.index.get_loc(i)
                break

    if insert_at is not None:
        top = rest_df.iloc[:insert_at]
        bottom = rest_df.iloc[insert_at:]
        merged = pd.concat([top, multi_df, bottom], ignore_index=True)
    else:
        merged = pd.concat([rest_df, multi_df], ignore_index=True)

    return merged


def main():
    xl = pd.ExcelFile(V4)
    df = pd.read_excel(V4, sheet_name="功能设计总表")
    others = {n: pd.read_excel(V4, sheet_name=n) for n in xl.sheet_names if n != "功能设计总表"}

    df = reorder_inference_deploy(df)

    # 阶段汇总（保持现有人月）
    summary = []
    for ph in ["阶段1", "阶段2", "阶段3"]:
        sub = df[df["阶段"] == ph]
        summary.append(
            {
                "阶段": ph,
                "功能点数": len(sub),
                "开发工作量(人月)": round(sub["开发工作量(人月)"].sum(), 1),
                "测试工作量(人月)": round(sub["测试工作量(人月)"].sum(), 1),
                "总工作量(人月)": round(sub["总工作量(人月)"].sum(), 1),
            }
        )

    full = df.copy()
    display = blank_hierarchy(full)

    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        display.to_excel(writer, sheet_name="功能设计总表", index=False)
        ws = writer.sheets["功能设计总表"]
        apply_styles(ws)
        merge_cells(ws, full)
        pd.DataFrame(summary).to_excel(writer, sheet_name="阶段汇总", index=False)
        apply_styles(writer.sheets["阶段汇总"])
        for name, sdf in others.items():
            if name != "阶段汇总":
                sdf.to_excel(writer, sheet_name=name, index=False)
                apply_styles(writer.sheets[name])

    multi = df[df["L3"].ffill().astype(str).str.contains("多机多卡", na=False)]
    print(f"Written: {OUTPUT}")
    print(f"多机多卡部署 L3 下 {len(multi)} 项 L4:")
    for _, r in multi.iterrows():
        print(f"  - {r['L4']}")


if __name__ == "__main__":
    main()
