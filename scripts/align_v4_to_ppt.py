#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Align AI平台功能设计矩阵_v4.xlsx with PPT targets: 91 items, 431 PM."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from format_v4_matrix import (
    apply_styles,
    blank_hierarchy,
    merge_cells,
)
from openpyxl import load_workbook

V4 = Path(r"C:\Users\d00804096\Desktop\AI平台功能设计矩阵_v4.xlsx")
V2 = Path(r"C:\Users\d00804096\Desktop\AI平台功能设计矩阵_v2.xlsx")

# PPT targets
PHASE_TARGETS = {"阶段1": 115.0, "阶段2": 276.0, "阶段3": 40.0}
L1_TARGETS = {
    "模型推理": 26,
    "模型训练": 16,
    "资产管理": 10,
    "算力调度": 12,
    "权限管理": 8,
    "指标监控": 9,
    "运维": 10,
}

# Remove 8 phase-2 items to reach 91 total (99 -> 91) and match L1 counts
REMOVE_L4 = [
    "多源数据接入",       # 资产管理 11 -> 10
    "公平共享DRF",        # 算力调度 15 -> 12
    "潮汐混部错峰",
    "超配与借用",
    "内容安全审核",       # 权限管理 9 -> 8
    "开发实例运维",       # 运维 13 -> 10
    "推理应用运维",
    "硬件兼容性检查",
]

PHASE_FINISH = {"阶段1": "2026年12月", "阶段2": "2027年6月", "阶段3": "2027年12月"}


def scale_phase(df: pd.DataFrame, phase: str, target: float) -> None:
    mask = df["阶段"] == phase
    cur = df.loc[mask, "总工作量(人月)"].sum()
    if cur <= 0:
        return
    factor = target / cur
    for col in ["开发工作量(人月)", "测试工作量(人月)", "总工作量(人月)"]:
        df.loc[mask, col] = (df.loc[mask, col] * factor).round(2)
    # fix rounding drift on last row of phase
    diff = target - df.loc[mask, "总工作量(人月)"].sum()
    if abs(diff) > 0.01:
        idx = df[mask].index[-1]
        df.at[idx, "总工作量(人月)"] = round(df.at[idx, "总工作量(人月)"] + diff, 2)
        dev_ratio = df.at[idx, "开发工作量(人月)"] / (
            df.at[idx, "开发工作量(人月)"] + df.at[idx, "测试工作量(人月)"] + 1e-9
        )
        df.at[idx, "开发工作量(人月)"] = round(df.at[idx, "总工作量(人月)"] * dev_ratio, 2)
        df.at[idx, "测试工作量(人月)"] = round(
            df.at[idx, "总工作量(人月)"] - df.at[idx, "开发工作量(人月)"], 2
        )


def build_l1_summary(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for l1, n in L1_TARGETS.items():
        sub = df[df["L1"] == l1]
        rows.append(
            {
                "L1领域": l1,
                "功能点数": len(sub),
                "目标功能点数": n,
                "开发工作量(人月)": round(sub["开发工作量(人月)"].sum(), 1),
                "测试工作量(人月)": round(sub["测试工作量(人月)"].sum(), 1),
                "总工作量(人月)": round(sub["总工作量(人月)"].sum(), 1),
            }
        )
    return pd.DataFrame(rows)


def main() -> None:
    xl = pd.ExcelFile(V4)
    df = pd.read_excel(V4, sheet_name="功能设计总表")
    other = {n: pd.read_excel(V4, sheet_name=n) for n in xl.sheet_names if n != "功能设计总表"}

    for c in ["L1", "L2", "L3"]:
        df[c] = df[c].ffill()

    before_n = len(df)
    before_pm = df["总工作量(人月)"].sum()

    # drop rows
    drop_mask = df["L4"].isin(REMOVE_L4)
    dropped = df[drop_mask]["L4"].tolist()
    df = df[~drop_mask].reset_index(drop=True)

    # fix MM mapping: MM has no heterogeneous capability
    het_mask = df["L4"] == "异构算力适配"
    if het_mask.any():
        df.loc[het_mask, "MM对应"] = ""

    # scale workloads per phase to PPT targets
    for ph, target in PHASE_TARGETS.items():
        scale_phase(df, ph, target)

    # summaries
    phase_rows = []
    for ph, target in PHASE_TARGETS.items():
        sub = df[df["阶段"] == ph]
        phase_rows.append(
            {
                "阶段": ph,
                "功能点数": len(sub),
                "目标功能点数": {**{"阶段1": 32, "阶段2": 54, "阶段3": 5}}[ph],
                "开发工作量(人月)": round(sub["开发工作量(人月)"].sum(), 1),
                "测试工作量(人月)": round(sub["测试工作量(人月)"].sum(), 1),
                "总工作量(人月)": round(sub["总工作量(人月)"].sum(), 1),
                "目标总工作量(人月)": target,
                "预计完成": PHASE_FINISH[ph],
            }
        )
    phase_summary = pd.DataFrame(phase_rows)
    l1_summary = build_l1_summary(df)

    width_ref = {}
    if V2.exists():
        wb2 = load_workbook(V2)
        ws2 = wb2["功能设计总表"]
        for col in range(1, ws2.max_column + 1):
            letter = ws2.cell(1, col).column_letter
            w = ws2.column_dimensions[letter].width
            if w:
                width_ref[letter] = w

    full = df.copy()
    display = blank_hierarchy(full)

    with pd.ExcelWriter(V4, engine="openpyxl") as writer:
        display.to_excel(writer, sheet_name="功能设计总表", index=False)
        ws = writer.sheets["功能设计总表"]
        apply_styles(ws, width_ref=width_ref or None)
        merge_cells(ws, full)
        phase_summary.to_excel(writer, sheet_name="阶段汇总", index=False)
        apply_styles(writer.sheets["阶段汇总"], width_ref=width_ref or None)
        l1_summary.to_excel(writer, sheet_name="L1汇总", index=False)
        apply_styles(writer.sheets["L1汇总"], width_ref=width_ref or None)
        for name, sdf in other.items():
            if name not in ("阶段汇总", "L1汇总"):
                sdf.to_excel(writer, sheet_name=name, index=False)
                if name in writer.sheets:
                    apply_styles(writer.sheets[name], width_ref=width_ref or None)

    print(f"Removed {before_n - len(df)} items: {dropped}")
    print(f"Items: {len(df)} (target 91)")
    print(f"Workload: {before_pm:.1f} -> {df['总工作量(人月)'].sum():.1f} (target 431)")
    print("\nPhase:")
    for ph in PHASE_TARGETS:
        sub = df[df["阶段"] == ph]
        print(f"  {ph}: {len(sub)} items, {sub['总工作量(人月)'].sum():.1f} PM")
    print("\nL1:")
    for l1 in L1_TARGETS:
        n = len(df[df["L1"] == l1])
        print(f"  {l1}: {n} (target {L1_TARGETS[l1]})")
    print(f"\nWritten: {V4}")


if __name__ == "__main__":
    main()
